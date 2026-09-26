"""Fichiers de l'utilisateur et exécution de code (Python, shell)."""
from __future__ import annotations

import asyncio
import base64
import io
import os
import signal
import sys
import time
import uuid
from pathlib import Path

from ..config import settings
from . import ToolContext, ToolResult, tool

IMAGE_EXT = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}


def read_any(path: Path, offset: int = 0, limit: int = 20000) -> tuple[str, list[dict]]:
    """Lit texte, PDF, Word, Excel, images. Renvoie (texte, images)."""
    ext = path.suffix.lower()
    if ext in IMAGE_EXT:
        from PIL import Image

        img = Image.open(path)
        img.thumbnail((1600, 1600))
        buf = io.BytesIO()
        img.convert("RGB").save(buf, "JPEG", quality=80)
        return f"Image {path.name} ({img.width}×{img.height})", [{"media_type": "image/jpeg", "data": base64.b64encode(buf.getvalue()).decode()}]
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n\n".join(f"--- page {i + 1} ---\n{p.extract_text() or ''}" for i, p in enumerate(reader.pages))
    elif ext == ".docx":
        import docx

        d = docx.Document(str(path))
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in t.rows)
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            rows = [";".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)]
            parts.append(f"## Feuille {ws.title}\n" + "\n".join(rows[:2000]))
        text = "\n\n".join(parts)
    else:
        raw = path.read_bytes()
        if b"\x00" in raw[:4096]:
            return f"Fichier binaire {path.name} ({len(raw)} octets) — utilise run_python pour le traiter.", []
        text = raw.decode("utf-8", errors="replace")
    total = len(text)
    chunk = text[offset: offset + limit]
    more = f"\n\n[… {total - offset - len(chunk)} caractères restants : relis avec offset={offset + len(chunk)}]" if offset + len(chunk) < total else ""
    return chunk + more, []


@tool("file_read", "Lit un fichier de l'espace de fichiers de l'utilisateur (texte, PDF, Word, Excel, image…).",
      {"path": {"type": "string"}, "offset": {"type": "integer"}, "limit": {"type": "integer"}},
      ["path"], label="Lecture de fichier", icon="📖", timeout=60)
async def file_read(ctx: ToolContext, path: str, offset: int = 0, limit: int = 20000) -> ToolResult:
    p = ctx.resolve_path(path)
    if not p.exists():
        return ToolResult(f"Fichier introuvable : {path}", is_error=True)
    if p.is_dir():
        return await file_list(ctx, path)
    text, images = await asyncio.to_thread(read_any, p, offset, limit)
    return ToolResult(text, images=images)


@tool("file_write", "Crée ou modifie un fichier texte dans l'espace de fichiers (l'utilisateur le voit et peut le télécharger). "
      "Fichiers intermédiaires (scripts, essais, brouillons) : dans .travail/, invisible pour lui.",
      {"path": {"type": "string"}, "content": {"type": "string"}, "append": {"type": "boolean"}},
      ["path", "content"], label="Écriture de fichier", icon="💾", timeout=30)
async def file_write(ctx: ToolContext, path: str, content: str, append: bool = False) -> ToolResult:
    p = ctx.resolve_path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "a" if append else "w", encoding="utf-8") as f:
        f.write(content)
    rel = ctx.rel(p)
    shown = not any(part.startswith(".") for part in Path(rel).parts)  # .travail/ & co : pas montré à l'utilisateur
    return ToolResult(f"Fichier {'complété' if append else 'écrit'} : {rel} ({p.stat().st_size} octets)", files=[rel] if shown else [])


@tool("file_list", "Liste les fichiers de l'espace de fichiers de l'utilisateur (documents reçus, créés, téléchargés).",
      {"path": {"type": "string", "description": "Sous-dossier (défaut : racine)"}}, [], label="Fichiers", icon="🗂️", timeout=30)
async def file_list(ctx: ToolContext, path: str = "") -> ToolResult:
    root = ctx.resolve_path(path or ".")
    if not root.exists():
        return ToolResult(f"Dossier introuvable : {path}", is_error=True)
    lines = []
    for p in sorted(root.rglob("*")):
        if any(part.startswith(".") for part in p.relative_to(root).parts) or len(lines) > 300:
            continue
        if p.is_file():
            st = p.stat()
            lines.append(f"{ctx.rel(p)}  ({st.st_size // 1024 or 1} Ko, {time.strftime('%d/%m/%Y %H:%M', time.localtime(st.st_mtime))})")
    return ToolResult("\n".join(lines) or "(dossier vide)")


def _snapshot(root: Path) -> dict[str, float]:
    out = {}
    for p in root.rglob("*"):
        if p.is_file() and ".ely" not in p.parts and ".sorties" not in p.parts:
            out[str(p)] = p.stat().st_mtime
    return out


async def _run(ctx: ToolContext, argv: list[str], timeout: int) -> ToolResult:
    ws = ctx.workspace
    before = _snapshot(ws)
    env = {**os.environ, "PYTHONIOENCODING": "utf-8", "MPLBACKEND": "Agg", "ELY_WORKSPACE": str(ws)}
    proc = await asyncio.create_subprocess_exec(*argv, cwd=str(ws), env=env, stdout=asyncio.subprocess.PIPE,
                                                stderr=asyncio.subprocess.STDOUT, start_new_session=True)

    async def kill_all() -> None:  # tout le groupe de processus, pas seulement l'enfant direct
        try:
            os.killpg(proc.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        await proc.wait()

    try:
        out, _ = await asyncio.wait_for(proc.communicate(), timeout=timeout)
    except asyncio.TimeoutError:
        await kill_all()
        return ToolResult(f"Arrêt après {timeout} s (délai dépassé).", is_error=True)
    except BaseException:  # tâche annulée ou arrêt d'Ely
        await asyncio.shield(kill_all())
        raise
    text = out.decode("utf-8", errors="replace").strip()
    after = _snapshot(ws)
    changed = [ctx.rel(Path(p)) for p, m in after.items() if before.get(p) != m]
    files_note = f"\nFichiers créés/modifiés : {', '.join(changed)}" if changed else ""
    status = "" if proc.returncode == 0 else f"\n(code de sortie {proc.returncode})"
    return ToolResult((text or "(aucune sortie)") + status + files_note, is_error=proc.returncode != 0, files=changed[:10])


@tool("run_python", """Exécute du code Python 3 (répertoire courant = espace de fichiers de l'utilisateur).
Bibliothèques disponibles : requests/httpx, pandas si installé, openpyxl, python-docx, reportlab, pypdf, pillow, numpy, matplotlib si installé.
Idéal pour calculs, analyse de données, conversion et génération de documents (Word, Excel, PDF, graphiques). Affiche les résultats avec print().""",
      {"code": {"type": "string"}, "timeout": {"type": "integer", "description": "secondes (défaut 120)"}},
      ["code"], label="Python", icon="🐍", timeout=900, available=lambda ctx: settings.allow_code)
async def run_python(ctx: ToolContext, code: str, timeout: int = 120) -> ToolResult:
    tmp = ctx.workspace / ".ely"
    tmp.mkdir(exist_ok=True)
    script = tmp / f"script_{uuid.uuid4().hex[:8]}.py"
    script.write_text(code, encoding="utf-8")
    try:
        return await _run(ctx, [sys.executable, str(script)], max(5, min(timeout, 880)))
    finally:
        script.unlink(missing_ok=True)


@tool("run_shell", "Exécute une commande shell (bash) sur le serveur d'Ely, dans l'espace de fichiers de l'utilisateur.",
      {"command": {"type": "string"}, "timeout": {"type": "integer", "description": "secondes (défaut 120)"}},
      ["command"], label="Terminal", icon="🖥️", timeout=900,
      available=lambda ctx: settings.allow_code and (ctx.is_admin or settings.allow_shell_for_all))
async def run_shell(ctx: ToolContext, command: str, timeout: int = 120) -> ToolResult:
    return await _run(ctx, ["bash", "-lc", command], max(5, min(timeout, 880)))
